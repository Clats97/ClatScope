from urllib.parse import quote
import requests
from pystyle import Colors, Write
from phonenumbers import geocoder, carrier
import phonenumbers
import os
import socket
import ssl
from concurrent.futures import ThreadPoolExecutor, as_completed
import dns.resolver
from dns import reversename
from email_validator import validate_email, EmailNotValidError
from urllib.parse import quote
import json
from bs4 import BeautifulSoup
import re
from email.parser import Parser
import whois
from tqdm import tqdm
from datetime import datetime
import magic
import stat
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS
import PyPDF2
import openpyxl
import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx import Presentation
from mutagen.easyid3 import EasyID3
from mutagen.mp3 import MP3
from mutagen.mp4 import MP4
from mutagen.id3 import ID3
from mutagen.flac import FLAC
import wave
from mutagen.oggvorbis import OggVorbis
from tinytag import TinyTag

_global_session = requests.Session()
default_color = Colors.light_red
requests.get = _global_session.get

import multiprocessing
MAX_WORKERS = min(32, (multiprocessing.cpu_count() or 1) * 5)

def validate_domain_input(domain):
    if not domain or len(domain) > 253 or ".." in domain:
        return False
    pattern = r"^[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$"
    return bool(re.match(pattern, domain))

def log_option(output_text):
    print()
    print("[?] Would you like to save this output to a log file? (Y/N): ", end="")
    choice = input().strip().upper()
    if choice == 'Y':
        stamp = datetime.now().strftime("[%Y-%m-%d %H:%M:%S] ")
        with open("clatscope_log.txt", "a", encoding="utf-8") as log_file:
            log_file.write(f"{stamp}{output_text}\n\n")
        Write.Print("[!] > Output has been saved to clatscope_log.txt\n", default_color, interval=0)

def export_json(data, filename_prefix="output"):
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{filename_prefix}_{stamp}.json"
    try:
        with open(filename, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)
        Write.Print(f"[!] > JSON Export complete: {filename}\n", Colors.green, interval=0)
    except Exception as e:
        Write.Print(f"[!] > Failed to write JSON file: {str(e)}\n", Colors.red, interval=0)

def clear():
    os.system('cls' if os.name == 'nt' else 'clear')

def restart():
    Write.Input("\nPress Enter to return to the main menu...", default_color, interval=0)
    clear()

def get_ip_details(ip):
    try:
        response = requests.get(f"https://ipinfo.io/{ip}/json", timeout=60)
        response.raise_for_status()
        return response.json()
    except:
        return None

def ip_info(ip):
    url = f"https://ipinfo.io/{ip}/json"
    try:
        response = requests.get(url)
        response.raise_for_status()
        data = response.json()
        loc = data.get('loc', 'None')
        maps_link = f"https://www.google.com/maps?q={loc}" if loc != 'None' else 'None'
        ip_details = f"""
╭─{' '*78}─╮
|{' '*34} IP Details {' '*34}|
|{'='*80}|
| [+] > IP Address         || {data.get('ip', 'None'):<51}|
| [+] > City               || {data.get('city', 'None'):<51}|
| [+] > Region             || {data.get('region', 'None'):<51}|
| [+] > Country            || {data.get('country', 'None'):<51}|
| [+] > Postal/ZIP Code    || {data.get('postal', 'None'):<51}|
| [+] > ISP                || {data.get('org', 'None'):<51}|
| [+] > Coordinates        || {loc:<51}|
| [+] > Timezone           || {data.get('timezone', 'None'):<51}|
| [+] > Location           || {maps_link:<51}|
╰─{' '*24}─╯╰─{' '*50}─╯
"""
        Write.Print(ip_details, Colors.white, interval=0)
        log_option(ip_details)

        print("[?] Export IP details in JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json(data, filename_prefix="ip_info")

    except:
        clear()
        Write.Print("\n[!] > Error retrieving IP address info.", default_color, interval=0)
    restart()

def subdomain_enumeration(domain):
    import requests
    from datetime import datetime
    url = f"https://crt.sh/?q=%.{domain}&output=json"
    Write.Print(f"\n[!] Subdomain enumeration for: {domain}\n", Colors.white, interval=0)
    try:
        resp = requests.get(url, timeout=60)
        if resp.status_code == 200:
            try:
                data = resp.json()
            except json.JSONDecodeError:
                msg = "[!] > Error: crt.sh returned non-JSON or empty data.\n"
                Write.Print(msg, Colors.red, interval=0)
                return
            found_subs = set()
            for entry in data:
                if 'name_value' in entry:
                    for subd in entry['name_value'].split('\n'):
                        subd_strip = subd.strip()
                        if subd_strip and subd_strip != domain:
                            found_subs.add(subd_strip)
                elif 'common_name' in entry:
                    c = entry['common_name'].strip()
                    if c and c != domain:
                        found_subs.add(c)
            if found_subs:
                out_text = f"\n[+] Found {len(found_subs)} subdomains for {domain}:\n"
                for s in sorted(found_subs):
                    out_text += f"    {s}\n"
                Write.Print(out_text, Colors.green, interval=0)
                print()
                print("[?] Would you like to save this output to a log file? (Y/N): ", end="")
                choice = input().strip().upper()
                if choice == 'Y':
                    stamp = datetime.now().strftime("[%Y-%m-%d %H:%M:%S] ")
                    with open("clatscope_log.txt", "a", encoding="utf-8") as f:
                        f.write(stamp + out_text + "\n")
                    Write.Print("[!] > Subdomains saved to clatscope_log.txt\n", Colors.white, interval=0)

                print("[?] Export subdomains as JSON? (Y/N): ", end="")
                if input().strip().upper() == "Y":
                    export_json(list(found_subs), filename_prefix="subdomains")

            else:
                Write.Print("[!] > No subdomains found.\n", Colors.red, interval=0)
        else:
            err = f"[!] > HTTP {resp.status_code} from crt.sh\n"
            Write.Print(err, Colors.red, interval=0)
    except Exception as exc:
        Write.Print(f"[!] > Subdomain enumeration error: {exc}\n", Colors.red, interval=0)

def deep_account_search(nickname):
    sites = [
        "https://youtube.com/@{target}",
        "https://facebook.com/{target}",
        "https://wikipedia.org/wiki/User:{target}",
        "https://instagram.com/{target}",
        "https://reddit.com/user/{target}",
        "https://medium.com/@{target}",
        "https://www.quora.com/profile/{target}",
        "https://bing.com/{target}",
        "https://x.com/{target}",
        "https://yandex.ru/{target}",
        "https://whatsapp.com/{target}",
        "https://yahoo.com/{target}",
        "https://amazon.com/{target}",
        "https://duckduckgo.com/{target}",
        "https://yahoo.co.jp/{target}",
        "https://tiktok.com/@{target}",
        "https://msn.com/{target}",
        "https://netflix.com/{target}",
        "https://weather.com/{target}",
        "https://live.com/{target}",
        "https://naver.com/{target}",
        "https://microsoft.com/{target}",
        "https://twitch.tv/{target}",
        "https://office.com/{target}",
        "https://vk.com/{target}",
        "https://pinterest.com/{target}",
        "https://discord.com/{target}",
        "https://aliexpress.com/{target}",
        "https://github.com/{target}",
        "https://adobe.com/{target}",
        "https://rakuten.co.jp/{target}",
        "https://ikea.com/{target}",
        "https://bbc.co.uk/{target}",
        "https://amazon.co.jp/{target}",
        "https://speedtest.net/{target}",
        "https://samsung.com/{target}",
        "https://healthline.com/{target}",
        "https://medlineplus.gov/{target}",
        "https://roblox.com/users/{target}/profile",
        "https://cookpad.com/{target}",
        "https://indiatimes.com/{target}",
        "https://mercadolivre.com.br/{target}",
        "https://britannica.com/{target}",
        "https://merriam-webster.com/{target}",
        "https://hurriyet.com.tr/{target}",
        "https://steamcommunity.com/user/{target}",
        "https://booking.com/{target}",
        "https://support.google.com/{target}",
        "https://bbc.com/{target}",
        "https://playstation.com/{target}",
        "https://ebay.com/usr/{target}",
        "https://poki.com/{target}",
        "https://walmart.com/{target}",
        "https://medicalnewstoday.com/{target}",
        "https://gov.uk/{target}",
        "https://nhs.uk/{target}",
        "https://detik.com/{target}",
        "https://cricbuzz.com/{target}",
        "https://nih.gov/{target}",
        "https://uol.com.br/{target}",
        "https://ilovepdf.com/{target}",
        "https://clevelandclinic.org/{target}",
        "https://cnn.com/{target}",
        "https://globo.com/{target}",
        "https://nytimes.com/{target}",
        "https://taboola.com/{target}",
        "https://pornhub.com/users/{target}",
        "https://redtube.com/users/{target}",
        "https://xnxx.com/profiles/{target}",
        "https://brazzers.com/profile/{target}",
        "https://xhamster.com/users/{target}",
        "https://onlyfans.com/{target}",
        "https://xvideos.es/profiles/{target}",
        "https://xvideos.com/profiles/{target}",
        "https://chaturbate.com/{target}",
        "https://redgifs.com/users/{target}",
        "https://tinder.com/{target}",
        "https://pof.com/{target}",
        "https://match.com/{target}",
        "https://eharmony.com/{target}",
        "https://bumble.com/{target}",
        "https://okcupid.com/{target}",
        "https://Badoo.com/{target}",
        "https://dating.com/{target}",
        "https://trello.com/{target}",
        "https://mapquest.com/{target}",
        "https://zoom.com/{target}",
        "https://apple.com/{target}",
        "https://dropbox.com/{target}",
        "https://weibo.com/{target}",
        "https://wordpress.com/{target}",
        "https://cloudflare.com/{target}",
        "https://salesforce.com/{target}",
        "https://fandom.com/{target}",
        "https://paypal.com/{target}",
        "https://soundcloud.com/{target}",
        "https://forbes.com/{target}",
        "https://theguardian.com/{target}",
        "https://hulu.com/{target}",
        "https://stackoverflow.com/users/{target}",
        "https://businessinsider.com/{target}",
        "https://huffpost.com/{target}",
        "https://booking.com/{target}",
        "https://pastebin.com/u/{target}",
        "https://producthunt.com/@{target}",
        "https://pypi.org/user/{target}",
        "https://slideshare.com/{target}",
        "https://strava.com/athletes/{target}",
        "https://tldrlegal.com/{target}",
        "https://t.me/{target}",
        "https://last.fm/user{target}",
        "https://data.typeracer.com/pit/profile?user={target}",
        "https://tryhackme.com/p/{target}",
        "https://trakt.tv/users/{target}",
        "https://scratch.mit.edu/users/{target}",
        "https://replit.com?{target}",
        "https://hackaday.io/{target}",
        "https://freesound.org/people/{target}",
        "https://hub.docker.com/u/{target}",
        "https://disqus.com/{target}",
        "https://www.codecademy.com/profiles/{target}",
        "https://www.chess.com/member/{target}",
        "https://bitbucket.org/{target}",
        "https://www.twitch.tv?{target}",
        "https://wikia.com/wiki/User:{target}",
        "https://steamcommunity.com/groups{target}",
        "https://keybase.io?{target}",
        "http://en.gravatar.com/{target}",
        "https://vk.com/{target}",
        "https://deviantart.com/{target}",
        "https://www.behance.net/{target}",
        "https://vimeo.com/{target}",
        "https://www.youporn.com/user/{target}",
        "https://profiles.wordpress.org/{target}",
        "https://tryhackme.com/p/{target}",
        "https://www.scribd.com/{target}",
        "https://myspace.com/{target}",
        "https://genius.com/{target}",
        "https://genius.com/artists/{target}",
        "https://www.flickr.com/people/{target}",
        "https://www.fandom.com/u/{target}",
        "https://www.chess.com/member/{target}",
        "https://buzzfeed.com/{target}",
        "https://www.buymeacoffee.com/{target}",
        "https://about.me/{target}",
        "https://discussions.apple.com/profile/{target}",
        "https://giphy.com/{target}",
        "https://scholar.harvard.edu/{target}",
        "https://www.instructables.com/member/{target}",
        "http://www.wikidot.com/user:info/{target}",
        "https://erome.com/{target}",
        "https://www.alik.cz/u/{target}",
        "https://rblx.trade/p/{target}",
        "https://www.paypal.com/paypalme/{target}",
        "https://hackaday.io/{target}",
        "https://connect.garmin.com/modern/profile/{target}"
    ]
    urls = [site_format.format(target=nickname) for site_format in sites]

    def check_url(url):
        try:
            response = requests.get(url, timeout=60)
            status_code = response.status_code
            if status_code == 200:
                return f"[+] > {url:<50}|| Found"
            elif status_code == 404:
                return f"[-] > {url:<50}|| Not found"
            else:
                return f"[-] > {url:<50}|| Error: {status_code}"
        except requests.exceptions.Timeout:
            return f"[-] > {url:<50}|| Timeout"
        except requests.exceptions.ConnectionError:
            return f"[-] > {url:<50}|| Connection error"
        except requests.exceptions.RequestException:
            return f"[-] > {url:<50}|| Request error"
        except Exception:
            return f"[-] > {url:<50}|| Unexpected error"

    title = "Deep Account Search"
    def fetch_social_urls(urls, title):
        result_str = f"""
╭─{' '*78}─╮
|{' '*27}{title}{' '*27}|
|{'='*80}|
"""
        with ThreadPoolExecutor() as executor:
            executor._max_workers = MAX_WORKERS
            results = list(executor.map(check_url, urls))
        for result in results:
            result_str += f"| {result:<78} |\n"
        result_str += f"╰─{' '*78}─╯"
        return result_str

    search_results = fetch_social_urls(urls, "Deep Account Search")
    Write.Print(search_results, Colors.white, interval=0)
    log_option(search_results)

    print("[?] Export deep account search to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({"nickname": nickname, "results": search_results}, filename_prefix="deep_account_search")
    restart()

def phone_info(phone_number):
    try:
        parsed_number = phonenumbers.parse(phone_number)
        country = geocoder.country_name_for_number(parsed_number, "en")
        region = geocoder.description_for_number(parsed_number, "en")
        operator = carrier.name_for_number(parsed_number, "en") if carrier else "" #the fix
        valid = phonenumbers.is_valid_number(parsed_number)
        validity = "Valid" if valid else "Invalid"
        phonetext = f"""
╭─{' '*50}─╮
|{' '*17}Phone number info{' '*18}|
|{'='*52}|
| [+] > Number   || {phone_number:<33}|
| [+] > Country  || {country:<33}     |
| [+] > Region   || {region:<33}      |
| [+] > Operator || {operator:<33}    |
| [+] > Validity || {validity:<33}    |
╰─{' '*15}─╯╰─{' '*31}─╯
"""
        Write.Print(phonetext, Colors.white, interval=0)
        log_option(phonetext)

        print("[?] Export phone info to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({
                "phone_number": phone_number,
                "country": country,
                "region": region,
                "operator": operator,
                "validity": validity
            }, filename_prefix="phone_info")

    except phonenumbers.phonenumberutil.NumberParseException:
        clear()
        Write.Print(f"\n[!] > Error: invalid phone number format (+1-000-000-0000)", default_color, interval=0)
    restart()

def dns_lookup(domain):
    record_types = ['A', 'CNAME', 'MX', 'NS']
    result_output = f"""
╭─{' '*78}─╮
|{' '*33} DNS Lookup {' '*33}|
|{'='*80}|
"""
    for rtype in record_types:
        result_output += f"| [+] > {rtype} Records: {' '*62}|\n"
        try:
            answers = dns.resolver.resolve(domain, rtype)
            for ans in answers:
                if rtype == 'MX':
                    result_output += f"|    {ans.preference:<4} {ans.exchange:<70}|\n"
                else:
                    result_output += f"|    {str(ans):<76}|\n"
        except dns.resolver.NoAnswer:
            result_output += "|    No records found.\n"
        except dns.resolver.NXDOMAIN:
            result_output += "|    Domain does not exist.\n"
        except Exception:
            result_output += "|    Error retrieving records.\n"
        result_output += f"|{'='*80}|\n"
    result_output += f"╰─{' '*78}─╯"
    Write.Print(result_output, Colors.white, interval=0)
    log_option(result_output)

    print("[?] Export DNS lookup to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({"domain": domain, "dns_records_raw": result_output}, filename_prefix="dns_lookup")

    restart()

def email_lookup(email_address):
    try:
        v = validate_email(email_address)
        email_domain = v.domain
    except EmailNotValidError as e:
        Write.Print(f"[!] > Invalid email address format: {str(e)}", default_color, interval=0)
        restart()
        return
    mx_records = []
    try:
        answers = dns.resolver.resolve(email_domain, 'MX')
        for rdata in answers:
            mx_records.append(str(rdata.exchange))
    except:
        mx_records = []
    validity = "Mx Found (Might be valid)" if mx_records else "No MX found (Might be invalid)"
    email_text = f"""
╭─{' '*78}─╮
|{' '*34}Email Info{' '*34}|
|{'='*80}|
| [+] > Email:        || {email_address:<52}|
| [+] > Domain:       || {email_domain:<52}|
| [+] > MX Records:   || {", ".join(mx_records) if mx_records else "None":<52}|
| [+] > Validity:     || {validity:<52}|
╰─{' '*23}─╯╰─{' '*51}─╯
"""
    Write.Print(email_text, Colors.white, interval=0)
    log_option(email_text)

    print("[?] Export email lookup to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({
            "email": email_address,
            "domain": email_domain,
            "mx_records": mx_records,
            "validity": validity
        }, filename_prefix="email_lookup")

    restart()

def reverse_dns(ip):
    try:
        rev_name = reversename.from_address(ip)
        answers = dns.resolver.resolve(rev_name, "PTR")
        ptr_record = str(answers[0]).strip('.')
    except:
        ptr_record = "No PTR record found"
    rdns_text = f"""
╭─{' '*78}─╮
|{' '*33}Reverse DNS Lookup{' '*33}|
|{'='*80}|
| [+] > IP:     || {ip:<60}|
| [+] > Host:   || {ptr_record:<60}|
╰─{' '*23}─╯╰─{' '*51}─╯
"""
    Write.Print(rdns_text, Colors.white, interval=0)
    log_option(rdns_text)

    print("[?] Export reverse DNS to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({"ip": ip, "ptr_record": ptr_record}, filename_prefix="reverse_dns")

    restart()

def analyze_email_header(raw_headers):
    parser = Parser()
    msg = parser.parsestr(raw_headers)
    from_ = msg.get("From", "")
    to_ = msg.get("To", "")
    subject_ = msg.get("Subject", "")
    date_ = msg.get("Date", "")
    received_lines = msg.get_all("Received", [])
    found_ips = []
    if received_lines:
        for line in received_lines:
            potential_ips = re.findall(r'\b\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}\b', line)
            for ip in potential_ips:
                if ip not in found_ips:
                    found_ips.append(ip)

    header_text = f"""
╭─{' '*78}─╮
|{' '*31}Email Header Analysis{' '*31}|
|{'='*80}|
| [+] > From:      || {from_:<55}|
| [+] > To:        || {to_:<55}|
| [+] > Subject:   || {subject_:<55}|
| [+] > Date:      || {date_:<55}|
|{'-'*80}|
"""
    if found_ips:
        header_text += "| [+] > Received Path (IPs found):\n"
        for ip in found_ips:
            header_text += f"|    {ip:<76}|\n"
    else:
        header_text += "| [+] > No IPs found in Received headers.\n"
    header_text += f"╰─{' '*78}─╯"
    Write.Print(header_text, Colors.white, interval=0)

    ip_details_full = ""
    if found_ips:
        ip_details_header = f"""
╭─{' '*78}─╮
|{' '*30}IP Geolocation Details{' '*30}|
|{'='*80}|
"""
        ip_details_summary = ""
        for ip in found_ips:
            data = get_ip_details(ip)
            if data is not None:
                loc = data.get('loc', 'None')
                ip_details_summary += f"| IP: {ip:<14}|| City: {data.get('city','N/A'):<15} Region: {data.get('region','N/A'):<15} Country: {data.get('country','N/A'):<4}|\n"
                ip_details_summary += f"|    Org: {data.get('org','N/A'):<63}|\n"
                ip_details_summary += f"|    Loc: {loc:<63}|\n"
                ip_details_summary += "|" + "-"*78 + "|\n"
            else:
                ip_details_summary += f"| IP: {ip:<14}|| [!] Could not retrieve details.\n"
                ip_details_summary += "|" + "-"*78 + "|\n"
        ip_details_footer = f"╰─{' '*78}─╯"
        ip_details_full = ip_details_header + ip_details_summary + ip_details_footer
        Write.Print(ip_details_full, Colors.white, interval=0)

    spf_result, dkim_result, dmarc_result = None, None, None
    spf_domain, dkim_domain = None, None
    auth_results = msg.get_all("Authentication-Results", [])
    from_domain = ""
    if "@" in from_:
        from_domain = from_.split("@")[-1].strip(">").strip()
    if auth_results:
        for entry in auth_results:
            spf_match = re.search(r'spf=(pass|fail|softfail|neutral)', entry, re.IGNORECASE)
            if spf_match:
                spf_result = spf_match.group(1)
            spf_domain_match = re.search(r'envelope-from=([^;\s]+)', entry, re.IGNORECASE)
            if spf_domain_match:
                spf_domain = spf_domain_match.group(1)
            dkim_match = re.search(r'dkim=(pass|fail|none|neutral)', entry, re.IGNORECASE)
            if dkim_match:
                dkim_result = dkim_match.group(1)
            dkim_domain_match = re.search(r'd=([^;\s]+)', entry, re.IGNORECASE)
            if dkim_domain_match:
                dkim_domain = dkim_domain_match.group(1)
            dmarc_match = re.search(r'dmarc=(pass|fail|none)', entry, re.IGNORECASE)
            if dmarc_match:
                dmarc_result = dmarc_match.group(1)
    spf_align = False
    dkim_align = False
    if from_domain and spf_domain:
        spf_align = from_domain.lower() == spf_domain.lower()
    if from_domain and dkim_domain:
        dkim_align = from_domain.lower() == dkim_domain.lower()
    alignment_text = f"""
╭─{' '*78}─╮
|{' '*30}SPF / DKIM / DMARC Checks{' '*29}|
|{'='*80}|
| [+] > SPF  Result:   {spf_result if spf_result else 'Not found':<20}   Domain: {spf_domain if spf_domain else 'N/A':<20} Aligned: {spf_align}|
| [+] > DKIM Result:   {dkim_result if dkim_result else 'Not found':<20} Domain: {dkim_domain if dkim_domain else 'N/A':<20} Aligned: {dkim_align}|
| [+] > DMARC Result:  {dmarc_result if dmarc_result else 'Not found':<20}|
╰─{' '*78}─╯
"""
    Write.Print(alignment_text, Colors.white, interval=0)
    full_output = header_text + "\n" + ip_details_full + "\n" + alignment_text
    log_option(full_output)

    print("[?] Export email header analysis to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({
            "raw_headers": raw_headers,
            "from": from_,
            "to": to_,
            "subject": subject_,
            "date": date_,
            "found_ips": found_ips,
            "spf_result": spf_result,
            "spf_domain": spf_domain,
            "spf_aligned": spf_align,
            "dkim_result": dkim_result,
            "dkim_domain": dkim_domain,
            "dkim_aligned": dkim_align,
            "dmarc_result": dmarc_result
        }, filename_prefix="email_header_analysis")
    restart()

def whois_lookup(domain):
    try:
        w = whois.whois(domain)
        clear()
        domain_name = w.domain_name if w.domain_name else "N/A"
        registrar = w.registrar if w.registrar else "N/A"
        creation_date = w.creation_date if w.creation_date else "N/A"
        expiration_date = w.expiration_date if w.expiration_date else "N/A"
        updated_date = w.updated_date if w.updated_date else "N/A"
        name_servers = ", ".join(w.name_servers) if w.name_servers else "N/A"
        status = ", ".join(w.status) if w.status else "N/A"
        whois_text = f"""
╭─{' '*78}─╮
|{' '*34}WHOIS Lookup{' '*34}|
|{'='*80}|
| [+] > Domain Name:       || {str(domain_name):<52}|
| [+] > Registrar:         || {str(registrar):<52}|
| [+] > Creation Date:     || {str(creation_date):<52}|
| [+] > Expiration Date:   || {str(expiration_date):<52}|
| [+] > Updated Date:      || {str(updated_date):<52}|
| [+] > Name Servers:      || {name_servers:<52}|
| [+] > Status:            || {status:<52}|
╰─{' '*23}─╯╰─{' '*51}─╯
"""
        Write.Print(whois_text, Colors.white, interval=0)
        log_option(whois_text)

        print("[?] Export WHOIS to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            data = {
                "domain": domain,
                "domain_name": str(domain_name),
                "registrar": str(registrar),
                "creation_date": str(creation_date),
                "expiration_date": str(expiration_date),
                "updated_date": str(updated_date),
                "name_servers": name_servers,
                "status": status
            }
            export_json(data, filename_prefix="whois_lookup")
    except Exception as e:
        clear()
        Write.Print(f"[!] > WHOIS lookup error: {str(e)}", default_color, interval=0)
    restart()

def check_password_strength(password):
    txt_file_path = os.path.join(os.path.dirname(__file__), "passwords.txt")
    if os.path.isfile(txt_file_path):
        try:
            with open(txt_file_path, "r", encoding="utf-8") as f:
                common_words = f.read().splitlines()
            for word in common_words:
                if word and word in password:
                    return "Weak password (may contain common phrase, term, word, sequence, etc, DO NOT use this password)"
        except Exception:
            pass
    score = 0
    if len(password) >= 8:
        score += 1
    if len(password) >= 12:
        score += 1
    if re.search(r'[A-Z]', password):
        score += 1
    if re.search(r'[a-z]', password):
        score += 1
    if re.search(r'\d', password):
        score += 1
    if re.search(r'[^a-zA-Z0-9]', password):
        score += 1
    if score <= 2:
        return "Weak password (may contain common phrase, term, word, sequence, etc, DO NOT use this password)"
    elif 3 <= score <= 4:
        return "Moderate password (room for improvement)"
    else:
        return "Strong password (suitable for high security apps / credentials)"

def password_strength_tool():
    clear()
    Write.Print("[!] > Enter password to evaluate strength:\n", default_color, interval=0)
    password = Write.Input("[?] >  ", default_color, interval=0)
    if not password:
        clear()
        Write.Print("[!] > Password cannot be empty. Please enter the password.\n", default_color, interval=0)
        restart()
        return
    strength = check_password_strength(password)
    clear()
    output_text = f"Password Strength: {strength}\n"
    Write.Print(output_text, Colors.white, interval=0)
    log_option(output_text)
    restart()

def username_check():
    clear()
    Write.Print("[!] > Conducting Username Check...\n", default_color, interval=0)
    username = Write.Input("[?] > Enter the username: ", default_color, interval=0).strip()
    if not username:
        clear()
        Write.Print("[!] > No username provided.\n", Colors.red, interval=0)
        restart()
        return

    def fetch_wmn_data():
        try:
            response = requests.get("https://raw.githubusercontent.com/WebBreacher/WhatsMyName/main/wmn-data.json", timeout=60)
            response.raise_for_status()
            return response.json()
        except:
            Write.Print("[!] > Failed to fetch data from WhatsMyName repository.\n", Colors.red, interval=0)
            return None

    data = fetch_wmn_data()
    if not data:
        restart()
        return

    headers = {
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/132.0.0.0 Safari/537.36"
    }
    sites = data["sites"]
    total_sites = len(sites)
    found_sites = []
    output_accumulated = ""

    def check_site(site, username, headers):
        site_name = site["name"]
        uri_check = site["uri_check"].format(account=username)
        try:
            res = requests.get(uri_check, headers=headers, timeout=60)
            estring_pos = site["e_string"] in res.text
            estring_neg = site["m_string"] in res.text
            if res.status_code == site["e_code"] and estring_pos and not estring_neg:
                return site_name, uri_check
        except:
            pass
        return None

    try:
        with ThreadPoolExecutor(max_workers=20) as executor:
            futures = {executor.submit(check_site, site, username, headers): site for site in sites}
            with tqdm(total=total_sites, desc="Checking sites") as pbar:
                for future in as_completed(futures):
                    try:
                        result = future.result()
                        if result:
                            site_name, uri_check = result
                            found_sites.append((site_name, uri_check))
                            found_str = f"[+] Found on: {site_name}\n[+] Profile URL: {uri_check}\n"
                            output_accumulated += found_str
                            Write.Print(found_str, Colors.green, interval=0)
                    except Exception:
                        pass
                    finally:
                        pbar.update(1)
        if found_sites:
            summary_str = f"\n[!] > Username found on {len(found_sites)} sites!\n"
            output_accumulated += summary_str
            Write.Print(summary_str, Colors.green, interval=0)

            generate_html_report(username, found_sites)
            report_msg = f"\n[!] > Report saved: username_check_report_{username}.html\n"
            output_accumulated += report_msg
            Write.Print(report_msg, Colors.green, interval=0)
        else:
            no_result_str = f"[!] > No results found for {username}.\n"
            output_accumulated += no_result_str
            Write.Print(no_result_str, Colors.red, interval=0)
    except Exception as e:
        err_str = f"[!] > An error occurred: {str(e)}\n"
        output_accumulated += err_str
        Write.Print(err_str, Colors.red, interval=0)

    log_option(output_accumulated)
    print("[?] Export username check to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({"username": username, "found_sites": found_sites}, filename_prefix="username_check")
    restart()

def generate_html_report(username, found_sites):
    html_content = f"""
<html>
<head>
    <title>Username Check Report for {username}</title>
    <style>
        body {{
            font-family: Arial, sans-serif;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
        }}
        th, td {{
            border: 1px solid #ddd;
            padding: 8px;
            text-align: left;
        }}
        th {{
            background-color: #f2f2f2;
        }}
    </style>
</head>
<body>
    <h1>Username Check Report for {username}</h1>
    <table>
        <tr>
            <th>Website Name</th>
            <th>Profile URL</th>
        </tr>
"""
    for site_name, uri_check in found_sites:
        html_content += f"""
        <tr>
            <td>{site_name}</td>
            <td><a href="{uri_check}" target="_blank">{uri_check}</a></td>
        </tr>"""
    html_content += """
    </table>
</body>
</html>"""
    with open(f"username_check_report_{username}.html", "w") as report_file:
        report_file.write(html_content)

def check_ssl_cert(domain):
    try:
        context = ssl.create_default_context()
        with socket.create_connection((domain, 443), timeout=60) as sock:
            with context.wrap_socket(sock, server_hostname=domain) as ssock:
                cert = ssock.getpeercert()
        subject = dict(x[0] for x in cert['subject'])
        issued_to = subject.get('commonName', 'N/A')
        issuer = dict(x[0] for x in cert['issuer'])
        issued_by = issuer.get('commonName', 'N/A')
        not_before = cert['notBefore']
        not_after = cert['notAfter']
        not_before_dt = datetime.strptime(not_before, "%b %d %H:%M:%S %Y %Z")
        not_after_dt = datetime.strptime(not_after, "%b %d %H:%M:%S %Y %Z")
        info_text = f"""
╭─{' '*78}─╮
|{' '*33}SSL Certificate Info{' '*32}|
|{'='*80}|
| [+] > Domain:       {domain:<58}|
| [+] > Issued To:    {issued_to:<58}|
| [+] > Issued By:    {issued_by:<58}|
| [+] > Valid From:   {str(not_before_dt):<58}|
| [+] > Valid Until:  {str(not_after_dt):<58}|
╰─{' '*78}─╯
"""
        Write.Print(info_text, Colors.white, interval=0)
        log_option(info_text)

        print("[?] Export SSL info to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({
                "domain": domain,
                "issued_to": issued_to,
                "issued_by": issued_by,
                "valid_from": str(not_before_dt),
                "valid_until": str(not_after_dt)
            }, filename_prefix="ssl_info")

    except ssl.SSLError as e:
        Write.Print(f"[!] > SSL Error: {str(e)}\n", Colors.red, interval=0)
    except socket.timeout:
        Write.Print("[!] > Connection timed out.\n", Colors.red, interval=0)
    except Exception as e:
        Write.Print(f"[!] > An error occurred retrieving SSL cert info: {str(e)}\n", Colors.red, interval=0)
    restart()

def check_robots_and_sitemap(domain):
    urls = [
        f"https://{domain}/robots.txt",
        f"https://{domain}/sitemap.xml"
    ]
    result_text = f"""
╭─{' '*78}─╮
|{' '*32}Site Discovery{' '*32}|
|{'='*80}|
| [+] > Domain:  {domain:<63}|
|{'-'*80}|
"""
    for resource_url in urls:
        try:
            resp = requests.get(resource_url, timeout=60)
            if resp.status_code == 200:
                lines = resp.text.split('\n')
                result_text += f"| Resource: {resource_url:<66}|\n"
                result_text += f"| Status: 200 (OK)\n"
                result_text += f"|{'-'*80}|\n"
                snippet = "\n".join(lines[:10])
                snippet_lines = snippet.split('\n')
                for sline in snippet_lines:
                    trunc = sline[:78]
                    result_text += f"| {trunc:<78}|\n"
                if len(lines) > 10:
                    result_text += "| ... (truncated)\n"
            else:
                result_text += f"| Resource: {resource_url:<66}|\n"
                result_text += f"| Status: {resp.status_code}\n"
            result_text += f"|{'='*80}|\n"
        except requests.exceptions.RequestException as e:
            result_text += f"| Resource: {resource_url}\n"
            result_text += f"| Error: {str(e)}\n"
            result_text += f"|{'='*80}|\n"
    result_text += f"╰─{' '*78}─╯"
    Write.Print(result_text, Colors.white, interval=0)
    log_option(result_text)

    print("[?] Export robots/sitemap to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({"domain": domain, "discovery": result_text}, filename_prefix="site_discovery")
    restart()

def check_dnsbl(ip_address):
    dnsbl_list = [
        "zen.spamhaus.org",
        "bl.spamcop.net",
        "dnsbl.sorbs.net",
        "b.barracudacentral.org"
    ]
    reversed_ip = ".".join(ip_address.split(".")[::-1])
    results = []
    for dnsbl in dnsbl_list:
        query_domain = f"{reversed_ip}.{dnsbl}"
        try:
            answers = dns.resolver.resolve(query_domain, 'A')
            for ans in answers:
                results.append((dnsbl, str(ans)))
        except dns.resolver.NXDOMAIN:
            pass
        except dns.resolver.NoAnswer:
            pass
        except Exception as e:
            results.append((dnsbl, f"Error: {str(e)}"))
    report = f"""
╭─{' '*78}─╮
|{' '*33}DNSBL Check{' '*34}|
|{'='*80}|
| [+] > IP: {ip_address:<67}|
|{'-'*80}|
"""
    if results:
        report += "| The IP is listed on the following DNSBL(s):\n"
        for dnsbl, answer in results:
            report += f"|   {dnsbl:<25} -> {answer:<45}|\n"
    else:
        report += "| The IP is NOT listed on the tested DNSBL(s).\n"
    report += f"╰─{' '*78}─╯"
    Write.Print(report, Colors.white, interval=0)
    log_option(report)

    print("[?] Export DNSBL check to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({"ip_address": ip_address, "dnsbl_results": results}, filename_prefix="dnsbl_check")
    restart()

def fetch_webpage_metadata(url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/132.0.0.0 Safari/537.36"
    }
    try:
        resp = requests.get(url, headers=headers, timeout=60)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "lxml")
        title_tag = soup.find("title")
        meta_desc = soup.find("meta", attrs={"name": "description"})
        meta_keyw = soup.find("meta", attrs={"name": "keywords"})
        title = title_tag.get_text(strip=True) if title_tag else "N/A"
        description = meta_desc["content"] if meta_desc and "content" in meta_desc.attrs else "N/A"
        keywords = meta_keyw["content"] if meta_keyw and "content" in meta_keyw.attrs else "N/A"
        result_text = f"""
╭─{' '*78}─╮
|{' '*31}Webpage Metadata{' '*31}|
|{'='*80}|
| [+] > URL:         {url:<58}|
| [+] > Title:       {title:<58}|
| [+] > Description: {description:<58}|
| [+] > Keywords:    {keywords:<58}|
╰─{' '*78}─╯
"""
        Write.Print(result_text, Colors.white, interval=0)
        log_option(result_text)

        print("[?] Export webpage metadata to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({
                "url": url,
                "title": title,
                "description": description,
                "keywords": keywords
            }, filename_prefix="webpage_metadata")
    except Exception as e:
        Write.Print(f"[!] > Error fetching metadata: {str(e)}\n", Colors.red, interval=0)
    restart()

def ship_info(mmsi):
    if not mmsi:
        clear()
        Write.Print("[!] > Please enter a valid MMSI number.\n", default_color, interval=0)
        restart()
        return
    url = f"https://api.facha.dev/v1/ship/{mmsi}"
    try:
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        data = response.json()
        output_text = f"Ship Info for MMSI {mmsi}:\n" + json.dumps(data, indent=2)
        clear()
        Write.Print(output_text, Colors.white, interval=0)
        log_option(output_text)
        print("[?] Export ship info to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({"mmsi": mmsi, "data": data}, filename_prefix="ship_info")
    except Exception as e:
        clear()
        Write.Print(f"[!] > Error retrieving ship info: {str(e)}\n", Colors.red, interval=0)
    restart()

def ship_radius(latitude, longitude, radius):

    if not latitude or not longitude or not radius:
        clear()
        Write.Print("[!] > Please enter valid latitude, longitude, and radius values.\n", default_color, interval=0)
        restart()
        return

    url = f"https://api.facha.dev/v1/ship/radius/{latitude}/{longitude}/{radius}"
    try:
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        data = response.json()
        output_text = (
            f"Ships within a radius of {radius} around ({latitude}, {longitude}):\n"
            + json.dumps(data, indent=2)
        )
        clear()
        Write.Print(output_text, Colors.white, interval=0)
        log_option(output_text)
        print("[?] Export ship radius data to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({
                "latitude": latitude,
                "longitude": longitude,
                "radius": radius,
                "data": data
            }, filename_prefix="ship_radius")
    except Exception as e:
        clear()
        Write.Print(f"[!] > Error retrieving ship radius info: {str(e)}\n", Colors.red, interval=0)
    restart()

def aircraft_live_range(lat, lon, range_value):

    if not lat or not lon or not range_value:
        clear()
        Write.Print("[!] > Please enter valid latitude, longitude, and range values.\n", default_color, interval=0)
        restart()
        return

    url = f"https://api.facha.dev/v1/aircraft/live/range/{lat}/{lon}/{range_value}"
    try:
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        data = response.json()
        output_text = (
            f"Live aircraft within a range of {range_value} around ({lat}, {lon}):\n"
            + json.dumps(data, indent=2)
        )
        clear()
        Write.Print(output_text, Colors.white, interval=0)
        log_option(output_text)
        print("[?] Export aircraft live range data to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({
                "latitude": lat,
                "longitude": lon,
                "range": range_value,
                "data": data
            }, filename_prefix="aircraft_live_range")
    except Exception as e:
        clear()
        Write.Print(f"[!] > Error retrieving aircraft live range info: {str(e)}\n", Colors.red, interval=0)
    restart()

def aircraft_live_callsign(callsign):
    if not callsign:
        clear()
        Write.Print("[!] > Please enter a valid callsign.\n", default_color, interval=0)
        restart()
        return
    url = f"https://api.facha.dev/v1/aircraft/live/callsign/{callsign}"
    try:
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        data = response.json()
        output_text = (
            f"Live aircraft info for callsign '{callsign}':\n" +
            json.dumps(data, indent=2)
        )
        clear()
        Write.Print(output_text, Colors.white, interval=0)
        log_option(output_text)
        print("[?] Export aircraft callsign data to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({
                "callsign": callsign,
                "data": data
            }, filename_prefix="aircraft_live_callsign")
    except Exception as e:
        clear()
        Write.Print(f"[!] > Error retrieving aircraft info for callsign '{callsign}': {str(e)}\n", Colors.red, interval=0)
    restart()

def read_file_metadata(file_path):
    clear()
    Write.Print(f"🐢 Checking File Data\n {file_path}", Colors.green, interval=0)

    def timeConvert(atime):
        from datetime import datetime
        dt = atime
        newtime = datetime.fromtimestamp(dt)
        return newtime.date()

    def sizeFormat(size):
        newsize = format(size/1024, ".2f")
        return newsize + " KB"

    try:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File {file_path} does not exist.")
        Dfile = os.stat(file_path)
        file_size = sizeFormat(Dfile.st_size)
        file_name = os.path.basename(file_path)

        max_length = 60
        file_creation_time = datetime.fromtimestamp(getattr(Dfile, 'st_birthtime', Dfile.st_ctime)).date()
        file_modification_time = timeConvert(Dfile.st_mtime)
        file_last_Access_Date = timeConvert(Dfile.st_atime)

        mime = magic.Magic(mime=True)
        file_type = mime.from_file(file_path)

        metaData_extra = []

        def get_permission_string(file_mode):
            permissions = [
                stat.S_IRUSR, stat.S_IWUSR, stat.S_IXUSR,
                stat.S_IRGRP, stat.S_IWGRP, stat.S_IXGRP,
                stat.S_IROTH, stat.S_IWOTH, stat.S_IXOTH
            ]
            labels = ['Owner', 'Group', 'Other']
            permission_descriptions = []
            for i, label in enumerate(labels):
                read = 'Yes' if file_mode & permissions[i * 3] else 'No'
                write = 'Yes' if file_mode & permissions[i * 3 + 1] else 'No'
                execute = 'Yes' if file_mode & permissions[i * 3 + 2] else 'No'
                description = f"{label} {{Read: {read}, Write: {write}, Execute: {execute}}}"
                permission_descriptions.append(description)
            return ', '.join(permission_descriptions)

        def gps_extract(exif_dict):
            gps_metadata = exif_dict['GPSInfo']
            lat_ref_num = 1 if gps_metadata['GPSLatitudeRef'] == 'N' else -1
            long_ref_num = 1 if gps_metadata['GPSLongitudeRef'] == 'E' else -1

            lat_list = [float(num) for num in gps_metadata['GPSLatitude']]
            long_list = [float(num) for num in gps_metadata['GPSLongitude']]

            lat_coordinate = (lat_list[0] + lat_list[1]/60 + lat_list[2]/3600) * lat_ref_num
            long_coordinate = (long_list[0] + long_list[1]/60 + long_list[2]/3600) * long_ref_num
            return (lat_coordinate, long_coordinate)

        permissions = get_permission_string(Dfile.st_mode)

        if file_type.startswith("image"):
            with Image.open(file_path) as img:
                metaData_extra.append(f"|{' '*32}Image MetaData{' '*32}|")
                metaData_extra.append(f"|{'-'*78}|")
                info_dict = {
                    "Filename": img.filename,
                    "Image Size": img.size,
                    "Image Height": img.height,
                    "Image Width": img.width,
                    "Image Format": img.format,
                    "Image Mode": img.mode
                }
                for label,value in info_dict.items():
                    metaData_extra.append(f"|  {str(label):<10}: ||  {str(value)[:max_length]:<60}|")
                if img.format == 'TIFF':
                    for tag_id, value in img.tag_v2.items():
                        tag_name = TAGS.get(tag_id, tag_id)
                        metaData_extra.append(f"|  {str(tag_name):<10}: ||  {str(value)[:max_length]:<60}|")
                elif file_path.endswith('.png'):
                    for key, value in img.info.items():
                        metaData_extra.append(f"|  {str(key):<10}: ||  {str(value)[:max_length]:<60}|")
                else:
                    imdata = img._getexif()
                    if imdata:
                        for tag_id in imdata:
                            tag = TAGS.get(tag_id, tag_id)
                            data = imdata.get(tag_id)
                            if tag == "GPSInfo":
                                gps = gps_extract(imdata)
                                metaData_extra.append(f"|  GPS Coordinates: ||  {gps}  |")
                                continue
                            if isinstance(data, bytes):
                                try:
                                    data = data.decode('utf-8', errors='ignore')
                                except UnicodeDecodeError:
                                    data = '<Unintelligible Data>'
                            metaData_extra.append(f"|  {str(tag):<10}: ||  {str(data)[:max_length]:<60}|")
                    else:
                        metaData_extra.append("No EXIF data found.")
        elif file_type == "application/pdf":
            with open(file_path, "rb") as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                pdf_data = pdf_reader.metadata
                metaData_extra.append(f"|{' '*32}PDF Metadata{' '*32}|")
                metaData_extra.append(f"|{'-'*78}|")
                if pdf_data:
                    for key, value in pdf_data.items():
                        metaData_extra.append(f"|  {str(key):<10}: || {str(value)[:max_length]:<60}|")
                    if pdf_reader.is_encrypted:
                        metaData_extra.append(f"|  Encrypted: || Yes      |")
                    else:
                        metaData_extra.append(f"|  Encrypted: || No      |")
                else:
                    metaData_extra.append("No PDF metadata found.")
        elif file_path.endswith(('.doc', '.docx')):
            doc = docx.Document(file_path)
            core_properties = doc.core_properties
            doc_metadata = f"""
|{' '*32}Document Properties{' '*32}
|{'='*78}|
| Title:            || {str(core_properties.title) :<60}           |
| Author:           || {str(core_properties.author) :<60}          |
| Subject:          || {str(core_properties.subject) :<60}         |
| Keywords:         || {str(core_properties.keywords) :<60}        |
| Last Modified By: || {str(core_properties.last_modified_by) :<60}|
| Created:          || {str(core_properties.created) :<60}         |
| Modified:         || {str(core_properties.modified) :<60}        |
| Category:         || {str(core_properties.category) :<60}        |
| Content Status:   || {str(core_properties.content_status) :<60}  |
| Version:          || {str(core_properties.version) :<60}         |
| Revision:         || {str(core_properties.revision) :<60}        |
| Comments:         || {str(core_properties.comments) :<60}        |
            """
            metaData_extra.append(doc_metadata)
        elif file_path.endswith(('.xlsx', '.xlsm')):
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            properties = workbook.properties
            excel_metadata = f"""
|{' '*32}Excel Document Properties{' '*32}|
|{'='*78}|
| Title:            || {str(properties.title) :<60}         |
| Author:           || {str(properties.creator) :<60}       |
| Keywords:         || {str(properties.keywords) :<60}      |
| Last Modified By: || {str(properties.lastModifiedBy) :<60}|
| Created:          || {str(properties.created) :<60}       |
| Modified:         || {str(properties.modified) :<60}      |
| Category:         || {str(properties.category) :<60}      |
| Description:      || {str(properties.description) :<60}   |
            """
            metaData_extra.append(excel_metadata)
        elif file_path.endswith(('.pptx', '.pptm')):
            try:
                presentation = Presentation(file_path)
                core_properties = presentation.core_properties
                pptx_metadata = f"""
|{' '*32}PowerPoint Document Properties{' '*31}|
|{'='*78}|
| Title:            || {str(core_properties.title) :<60}           |
| Author:           || {str(core_properties.author) :<60}          |
| Keywords:         || {str(core_properties.keywords) :<60}        |
| Last Modified By: || {str(core_properties.last_modified_by) :<60}|
| Created:          || {str(core_properties.created) :<60}         |
| Modified:         || {str(core_properties.modified) :<60}        |
| Category:         || {str(core_properties.category) :<60}        |
| Description:      || {str(core_properties.subject) :<60}         |
                """
                metaData_extra.append(pptx_metadata)
            except Exception as e:
                metaData_extra.append(f"[Error] Could not read PowerPoint metadata: {e}")
        elif file_type.startswith("audio"):
            try:
                metaData_extra.append(f"|{' '*32}Audio MetaData{' '*32}|\n|{'-'*78}|")
                tinytim = TinyTag.get(file_path)
                if tinytim:
                    metaData_extra.append(f"|  Title:    || {str(tinytim.title)[:max_length]:<60}      |")
                    metaData_extra.append(f"|  Artist:   || {str(tinytim.artist)[:max_length]:<60}     |")
                    metaData_extra.append(f"|  Genre:    || {str(tinytim.genre)[:max_length]:<60}      |")
                    metaData_extra.append(f"|  Album:    || {str(tinytim.album)[:max_length]:<60}      |")
                    metaData_extra.append(f"|  Year:     || {str(tinytim.year)[:max_length]:<60}       |")
                    metaData_extra.append(f"|  Composer: || {str(tinytim.composer)[:max_length]:<60}   |")
                    metaData_extra.append(f"|  A-Artist: || {str(tinytim.albumartist)[:max_length]:<60}|")
                    metaData_extra.append(f"|  Track     || {str(tinytim.track_total)[:max_length]:<60}|")
                    metaData_extra.append(f"|  Duration: || {f'{tinytim.duration:.2f} seconds':<60}    |")
                    metaData_extra.append(f"|  Bitrate:  || {str(tinytim.bitrate) + ' kbps':<60}       |")
                    metaData_extra.append(f"|  Samplrate:|| {str(tinytim.samplerate) + ' Hz':<60}      |")
                    metaData_extra.append(f"|  Channels: || {str(tinytim.channels):<60}                |")

                if file_path.endswith('.mp3'):
                    audio = MP3(file_path, ID3=ID3)
                elif file_path.endswith('.wav'):
                    audio = wave.open(file_path, 'rb')
                elif file_path.endswith('.flac'):
                    audio = FLAC(file_path)
                elif file_path.endswith('.ogg'):
                    audio = OggVorbis(file_path)
                elif file_path.endswith(('.m4a', '.mp4')):
                    audio = MP4(file_path)
                else:
                    audio = None

                if audio is None:
                    metaData_extra.append(" 🐸 Cant Read Audio File for metadata.\n Unsupported")
                else:
                    if hasattr(audio, 'items') and audio.items():
                        for tag, value in audio.items():
                            metaData_extra.append(f"|  {str(tag):<10}: ||  {str(value)[:max_length]:<60}|")
            except Exception as e:
                metaData_extra.append(f"Error processing file: {str(e)}")

        clear()
        metadata_summary = f"""
|{' '*32}File Metadata{' '*33}|
|{'='*78}|
|  File Path:   || {file_path:<60}                  |
|  File Name:   || {file_name:<60}                  |
|  File Size:   || {file_size:<60}                  |
|  File Type:   || {file_type:<60}                  |
|  Permission:  || {permissions:<60}                |
|  Created:     || {str(file_creation_time):<60}    |
|  Modified:    || {str(file_modification_time):60}|
|  Last Access: || {str(file_last_Access_Date):60}  |
"""
        metadata_summary += "\n".join(metaData_extra)
        metadata_summary += "\n" + "="*78 + "\n"
        Write.Print(metadata_summary, Colors.white, interval=0)
        log_option(metadata_summary)

        print("[?] Export file metadata to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({
                "file_path": file_path,
                "file_name": file_name,
                "file_size": file_size,
                "file_type": file_type,
                "permissions": permissions,
                "created": str(file_creation_time),
                "modified": str(file_modification_time),
                "last_access": str(file_last_Access_Date),
                "additional_metadata": metaData_extra
            }, filename_prefix="file_metadata")

    except Exception as e:
        err_msg = f" ☠️ Error reading file metadata: {str(e)}"
        Write.Print(err_msg, Colors.red, interval=0)
        log_option(err_msg)
    restart()

def basic_port_scan(target, ports=[20, 21, 22, 80, 443, 8080, 23, 25, 53, 67, 68, 69, 88, 110, 123, 137, 138, 139, 143, 162, 162, 389, 427, 445, 465, 500, 636, 993, 995, 1433, 1434, 3306, 5060, 5061]):
    clear()
    Write.Print(f"[!] >Port Scan for {target}\n", default_color, interval=0)
    result_lines = []
    for port in ports:
        sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        sock.settimeout(1.0)
        try:
            result = sock.connect_ex((target, port))
            if result == 0:
                result_lines.append(f"Port {port} is OPEN.")
            else:
                result_lines.append(f"Port {port} is closed or filtered.")
        except Exception as e:
            result_lines.append(f"Port {port} -> Error: {str(e)}")
        finally:
            sock.close()

    report = f"\nPort Scan Results for {target}:\n" + "\n".join(result_lines)
    Write.Print(report, Colors.white, interval=0)
    log_option(report)

    print("[?] Export port scan results to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json({"target": target, "scan_results": result_lines}, filename_prefix="port_scan")
    restart()

def wayback_lookup(domain):
    clear()
    if not domain:
        Write.Print("[!] No domain provided for Wayback lookup.\n", Colors.red, interval=0)
        restart()
        return
    base_url = "http://web.archive.org/cdx/search/cdx"
    params = {
        "url": domain,
        "output": "json",
        "fl": "original,timestamp",
        "collapse": "digest",
        "filter": "statuscode:200",
        "limit": 20
    }
    try:
        resp = requests.get(base_url, params=params, timeout=60)
        resp.raise_for_status()
        data = resp.json()
        if len(data) <= 1:
            Write.Print("[!] > No historical snapshots found (or none with status 200).\n", Colors.red, interval=0)
            restart()
            return

        snapshots = data[1:]
        output_text = f"\nWayback Machine Snapshots for {domain}:\n"
        for snap in snapshots:
            original_url, timestamp = snap
            archive_url = f"https://web.archive.org/web/{timestamp}/{original_url}"
            output_text += f"- {timestamp} -> {archive_url}\n"

        Write.Print(output_text, Colors.white, interval=0)
        log_option(output_text)

        print("[?] Export Wayback data to JSON? (Y/N): ", end="")
        if input().strip().upper() == "Y":
            export_json({"domain": domain, "snapshots": snapshots}, filename_prefix="wayback_lookup")
    except Exception as e:
        Write.Print(f"[!] > Error fetching Wayback Machine data: {str(e)}\n", Colors.red, interval=0)
    restart()

def bulk_domain_processing(csv_path):
    clear()
    if not os.path.isfile(csv_path):
        Write.Print("[!] CSV file not found or invalid path.\n", Colors.red, interval=0)
        restart()
        return

    try:
        with open(csv_path, "r", encoding="utf-8") as f:
            lines = [x.strip() for x in f if x.strip()]
    except UnicodeDecodeError:
        with open(csv_path, "r", encoding="latin-1") as f:
            lines = [x.strip() for x in f if x.strip()]

    Write.Print(f"[!] Found {len(lines)} entries in {csv_path}.\n", default_color, interval=0)
    Write.Print("Choose the checks you want to run on each domain/IP:\n", Colors.white, interval=0)
    Write.Print("[1] DNS Lookup\n[2] WHOIS Lookup\n[3] Subdomain Enumeration\n[4] IP Info (if IP)\n[5] SSL Cert Info (if domain)\n\n", Colors.white, interval=0)
    chosen = Write.Input("[?] Enter your choices separated by commas (e.g. 1,2,3): ", default_color, interval=0).strip()
    chosen_set = set(x.strip() for x in chosen.split(","))

    results = {}
    for entry in lines:
        out_lines = [f"Results for {entry}:"]
        if "1" in chosen_set:
            try:
                out_lines.append("DNS Lookup:")
                try:
                    answers = dns.resolver.resolve(entry, 'A')
                    out_lines.append(f"A Records: {[str(a) for a in answers]}")
                except:
                    out_lines.append("No A records / Error retrieving.")
            except Exception as e:
                out_lines.append(f"DNS Lookup Error: {str(e)}")
        if "2" in chosen_set:
            try:
                w = whois.whois(entry)
                out_lines.append(f"WHOIS: {w.domain_name}, Registrar: {w.registrar}")
            except Exception as e:
                out_lines.append(f"WHOIS Error: {str(e)}")
        if "3" in chosen_set and validate_domain_input(entry):
            try:
                out_lines.append("Subdomain Enumeration: See script logs for details.")
            except Exception as e:
                out_lines.append(f"Subdomain enumeration error: {str(e)}")
        if "4" in chosen_set:
            if re.match(r'^\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}$', entry):
                details = get_ip_details(entry)
                if details:
                    out_lines.append(f"IP Info: City={details.get('city','N/A')} Region={details.get('region','N/A')} Org={details.get('org','N/A')}")
                else:
                    out_lines.append("IP Info not available.")
            else:
                out_lines.append("Not an IP, skipping IP Info.")
        if "5" in chosen_set and validate_domain_input(entry):
            try:
                context = ssl.create_default_context()
                with socket.create_connection((entry, 443), timeout=5) as sock:
                    with context.wrap_socket(sock, server_hostname=entry) as ssock:
                        cert = ssock.getpeercert()
                subject = dict(x[0] for x in cert['subject'])
                issuer = dict(x[0] for x in cert['issuer'])
                out_lines.append(f"SSL Issued To: {subject.get('commonName','N/A')}, By: {issuer.get('commonName','N/A')}")
            except Exception as e:
                out_lines.append(f"SSL Error: {str(e)}")

        results[entry] = "\n".join(out_lines)

    for entry, data in results.items():
        Write.Print("\n" + data + "\n", Colors.white, interval=0)

    print("[?] Export bulk domain results to JSON? (Y/N): ", end="")
    if input().strip().upper() == "Y":
        export_json(results, filename_prefix="bulk_domain")
    restart()

def main():
    while True:
        try:
            clear()
            print("\033[1;31m ██████╗██╗      █████╗ ████████╗███████╗ ██████╗ ██████╗ ██████╗ ███████╗")
            print("██╔════╝██║     ██╔══██╗╚══██╔══╝██╔════╝██╔════╝██╔═══██╗██╔══██╗██╔════╝")
            print("██║     ██║     ███████║   ██║   ███████╗██║     ██║   ██║██████╔╝█████╗  ")
            print("██║     ██║     ██╔══██║   ██║   ╚════██║██║     ██║   ██║██╔═══╝ ██╔══╝  ")
            print("╚██████╗███████╗██║  ██║   ██║   ███████║╚██████╗╚██████╔╝██║     ███████╗")
            print(" ╚═════╝╚══════╝╚═╝  ╚═╝   ╚═╝   ╚══════╝ ╚═════╝ ╚═════╝ ╚═╝     ╚══════╝\033[0m")
            print("\033[1;31m                        ███╗   ███╗██╗███╗   ██╗██╗                        ")
            print("                        ████╗ ████║██║████╗  ██║██║                        ")
            print("                        ██╔████╔██║██║██╔██╗ ██║██║                        ")
            print("                        ██║╚██╔╝██║██║██║╚██╗██║██║                        ")
            print("                        ██║ ╚═╝ ██║██║██║ ╚████║██║                        ")
            print("                        ╚═╝     ╚═╝╚═╝╚═╝  ╚═══╝╚═╝                       \033[0m")
            print("\033[1;34mC L A T S C O P E   I N F O       T O O L    M I N I\033[0m   \033[1;31m(Version 1.00)\033[0m")
            author = "🛡️ By Joshua M Clatney (Clats97) - Ethical Pentesting Enthusiast 🛡️"
            Write.Print(author + "\n[OSINT]\nOpen Sources. Clear Conclusions\n", Colors.white, interval=0)
            menu = (
                "==============================================================================================================|\n"
                "|  №   ||         Function          ||                          Description                                   |\n"
                "|======||===========================||========================================================================|\n"
                "| [1]  || IP Address Search         || Retrieves IP address info                                              |\n"
                "| [2]  || Deep Account Search       || Retrieves profiles from various websites                               |\n"
                "| [3]  || Phone Search              || Retrieves phone number info                                            |\n"
                "| [4]  || DNS Record Search         || Retrieves DNS records (A, CNAME, MX, NS)                               |\n"
                "| [5]  || Email MX Search           || Retrieves MX info for an email                                         |\n"
                "| [6]  || Reverse DNS Search        || Retrieves PTR records for an IP address                                |\n"
                "| [7]  || Email Header Search       || Retrieves info from an email header                                    |\n"
                "| [8]  || WHOIS Search              || Retrieves domain registration data                                     |\n"
                "| [9]  || Password Analyzer         || Retrieves password strength rating                                     |\n"
                "| [10] || Username Search           || Retrieves usernames from online accounts                               |\n"
                "| [11] || SSL Search                || Retrieves basic SSL certificate details from a URL                     |\n"
                "| [12] || Web Crawler Search        || Retrieves Robots.txt & Sitemap.xml file info                           |\n"
                "| [13] || DNSBL Search              || Retrieves IP DNS blacklist info                                        |\n"
                "| [14] || Web Metadata Search       || Retrieves meta tags and more from a webpage                            |\n"
                "| [15] || File Metadata Search      || Retrieves metadata from various file types                             |\n"
                "| [16] || Subdomain Search          || Retrieves subdomain info                                               |\n"
                "| [17] || Wayback Search            || Retrieves historical snapshots from the Wayback Machine                |\n"
                "| [18] || Port Scan Search          || Retrieves scan results on common ports                                 |\n"
                "| [19] || Bulk CSV Search           || Retrieves multiple checks in bulk from a CSV for domain/IP             |\n"
                "| [20] || Ship Search v1            || Retrieves ship data by searching an MMSI                               |\n"
                "| [21] || Ship Search v2            || Retrieves ship data by searching a location (via radius)               |\n"
                "| [22] || Aircraft Search v1        || Retrieves aircraft data by searching a location                        |\n"
                "| [23] || Aircraft Search v2        || Retrieves aircraft data by searching a callsign                        |\n"
                "| [0]  || Exit                      || Exit ClatScope Info Tool                                               |\n"
                "╰─    ─╯╰─                         ─╯╰─                                                                      ─╯\n"
            )

            Write.Print(menu, Colors.white, interval=0)
            choice = Write.Input("[?] >  ", default_color, interval=0).strip()
            if choice == "1":
                clear()
                ip = Write.Input("[?] > IP-Address: ", default_color, interval=0)
                if not ip:
                    clear()
                    Write.Print("[!] > Enter an IP Address\n", default_color, interval=0)
                    continue
                ip_info(ip)
            elif choice == "2":
                clear()
                nickname = Write.Input("[?] > Username: ", default_color, interval=0)
                if not nickname:
                    clear()
                    Write.Print("[!] > Enter the username\n", default_color, interval=0)
                    continue
                deep_account_search(nickname)
            elif choice == "3":
                clear()
                phone_number = Write.Input("[?] > Phone number: ", default_color, interval=0)
                if not phone_number:
                    clear()
                    Write.Print("[!] > Enter the phone number\n", default_color, interval=0)
                    continue
                phone_info(phone_number)
            elif choice == "4":
                clear()
                domain = Write.Input("[?] > Domain / URL: ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain / URL\n", default_color, interval=0)
                    continue
                dns_lookup(domain)
            elif choice == "5":
                clear()
                email = Write.Input("[?] > Email: ", default_color, interval=0)
                if not email:
                    clear()
                    Write.Print("[!] > Enter email\n", default_color, interval=0)
                    continue
                email_lookup(email)
            elif choice == "6":
                clear()
                ip = Write.Input("[?] > Enter an IP Address for a Reverse DNS Search: ", default_color, interval=0)
                if not ip:
                    clear()
                    Write.Print("[!] > Enter an IP address\n", default_color, interval=0)
                    continue
                reverse_dns(ip)
            elif choice == "7":
                clear()
                Write.Print("[!] > Paste the raw email headers below as one single string (end with empty line):\n", default_color, interval=0)
                lines = []
                while True:
                    line = input()
                    if not line.strip():
                        break
                    lines.append(line)
                raw_headers = "\n".join(lines)
                if not raw_headers.strip():
                    clear()
                    Write.Print("[!] > No email header was provided.\n", default_color, interval=0)
                    continue
                analyze_email_header(raw_headers)
            elif choice == "8":
                clear()
                domain = Write.Input("[?] > Enter a domain / URL for WHOIS lookup: ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain / URL\n", default_color, interval=0)
                    continue
                whois_lookup(domain)
            elif choice == "9":
                clear()
                password_strength_tool()
            elif choice == "10":
                clear()
                username_check()
            elif choice == "11":
                clear()
                domain = Write.Input("[?] > Enter a domain / URL for SSL certificate verification: ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain or URL\n", default_color, interval=0)
                    continue
                check_ssl_cert(domain)
            elif choice == "12":
                clear()
                domain = Write.Input("[?] > Enter domain to check for Robots.txt & Sitemap.xml file(s): ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain / URL\n", default_color, interval=0)
                    continue
                check_robots_and_sitemap(domain)
            elif choice == "13":
                clear()
                ip_address = Write.Input("[?] > Enter IP address to check DNSBL: ", default_color, interval=0)
                if not ip_address:
                    clear()
                    Write.Print("[!] > Enter an IP address\n", default_color, interval=0)
                    continue
                check_dnsbl(ip_address)
            elif choice == "14":
                clear()
                url = Write.Input("[?] > Enter URL for metadata extraction: ", Colors.white, interval=0)
                if not url:
                    clear()
                    Write.Print("[!] > Enter a URL\n", default_color, interval=0)
                    continue
                fetch_webpage_metadata(url)
            elif choice == "15":
                clear()
                file_path = Write.Input(" 🐸 Enter path to the file you want analyzed: ", default_color, interval=0)
                read_file_metadata(file_path)
            elif choice == "16":
                clear()
                domain = Write.Input("[?] > Enter domain for subdomain enumeration: ", default_color, interval=0)
                subdomain_enumeration(domain)
            elif choice == "17":
                clear()
                domain = Write.Input("[?] > Enter domain for Wayback lookup: ", default_color, interval=0)
                wayback_lookup(domain)
            elif choice == "18":
                clear()
                target = Write.Input("[?] > Enter IP or domain for port scan: ", default_color, interval=0)
                basic_port_scan(target)
            elif choice == "19":
                clear()
                csv_path = Write.Input("[?] > Enter path to CSV file: ", Colors.white, interval=0)
                bulk_domain_processing(csv_path)
            elif choice == "20":
                clear()
                mmsi = Write.Input("[?] > Enter the MMSI for ship lookup: ", default_color, interval=0).strip()
                if not mmsi:
                    clear()
                    Write.Print("[!] > Please enter an MMSI number.\n", default_color, interval=0)
                    continue
                ship_info(mmsi)
            elif choice == "21":
                clear()
                latitude = Write.Input("[?] > Enter latitude: ", default_color, interval=0).strip()
                longitude = Write.Input("[?] > Enter longitude: ", default_color, interval=0).strip()
                radius = Write.Input("[?] > Enter search radius: ", default_color, interval=0).strip()
                if not latitude or not longitude or not radius:
                    clear()
                    Write.Print("[!] > Please enter latitude, longitude, and radius.\n", default_color, interval=0)
                    continue
                ship_radius(latitude, longitude, radius)
            elif choice == "22":
                clear()
                lat = Write.Input("[?] > Enter latitude: ", default_color, interval=0).strip()
                lon = Write.Input("[?] > Enter longitude: ", default_color, interval=0).strip()
                range_value = Write.Input("[?] > Enter search range: ", default_color, interval=0).strip()
                if not lat or not lon or not range_value:
                    clear()
                    Write.Print("[!] > Please enter latitude, longitude, and range.\n", default_color, interval=0)
                    continue
                aircraft_live_range(lat, lon, range_value)
            elif choice == "23":
                clear()
                callsign = Write.Input("[?] > Enter aircraft callsign: ", default_color, interval=0).strip()
                if not callsign:
                    clear()
                    Write.Print("[!] > Please enter a callsign.\n", default_color, interval=0)
                    continue
                aircraft_live_callsign(callsign)
            elif choice == "0":
                clear()
                Write.Print("[!] > Exiting ClatScope Info Tool.\n", Colors.white, interval=0)
                break
            else:
                clear()
                Write.Print("[!] > Invalid input. Please try again.\n", Colors.white, interval=0)
        except KeyboardInterrupt:
            clear()
            Write.Print("[!] > Exiting on user request...\n", Colors.white, interval=0)
            break

if __name__ == "__main__":
    main()